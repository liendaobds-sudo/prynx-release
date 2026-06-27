"""Data_Source_Reader — lõi thuần đọc & chuẩn hoá nguồn dữ liệu cho VDP.

Module này KHÔNG import ReportLab/PDF nên là pure logic, dễ kiểm thử bằng
property-based testing (Hypothesis).

Task 2.1 hiện thực phần lõi:
- ``RecordTable``       : cấu trúc bảng record chuẩn hoá dùng chung mọi định dạng.
- ``DataSourceError``   : ngoại lệ mang mã lỗi + thông báo tiếng Việt.
- ``detect_encoding``   : nhận diện encoding trong {utf-8, utf-8-sig, windows-1258}.
- ``detect_delimiter``  : nhận diện delimiter trong {',', ';', '\\t'}.
- ``normalize_columns`` : khử trùng tên cột + ô tiêu đề rỗng (không mất cột).

Task 2.5 bổ sung đường CSV:
- ``parse_delimited``   : phân tích văn bản phân tách → ``RecordTable`` (dòng tiêu
  đề = dòng KHÔNG rỗng đầu tiên, bỏ qua dòng hoàn toàn rỗng).
- ``read_source``       : điều phối theo định dạng nguồn (CSV + XLSX hiện thực;
  Google Sheets để task 2.12).

Task 2.8 bổ sung đường Excel:
- ``read_xlsx``         : đọc ``.xlsx`` (openpyxl read_only) → ``RecordTable``;
  ô gộp gán giá trị về ô trên-trái, các ô còn lại để rỗng (Req 1.12).
- ``list_xlsx_sheets``  : liệt kê tên sheet để người dùng chọn (Req 1.3).

Task 2.12 bổ sung đường Google Sheets:
- ``fetch_gsheet_csv``  : chuyển link Google Sheets thành URL export CSV và tải
  về bằng ``httpx``; HTTP 4xx hoặc redirect tới trang đăng nhập →
  ``GSHEET_FORBIDDEN`` (Req 1.4, 1.13). Được nối vào ``read_source`` (nhánh
  ``gsheet``) rồi đi qua đường CSV chung (:func:`_read_csv`).
"""

from __future__ import annotations

import codecs
import csv
import io
import re
from dataclasses import dataclass, field
from typing import Dict, List, Optional, Union
from urllib.parse import parse_qs, urlparse

# ─── Hằng số cấu hình nhận diện ──────────────────────────────────────────────

# Tập encoding hỗ trợ, thử theo đúng thứ tự ưu tiên (Req 1.6).
# UTF-8 BOM trước (nhận diện qua BOM), rồi UTF-8 strict, cuối cùng Windows-1258.
SUPPORTED_ENCODINGS: tuple[str, ...] = ("utf-8-sig", "utf-8", "windows-1258")

# Tập delimiter hỗ trợ cho file văn bản phân tách (Req 1.5).
SUPPORTED_DELIMITERS: tuple[str, ...] = (",", ";", "\t")


# ─── Cấu trúc dữ liệu chuẩn hoá ──────────────────────────────────────────────


@dataclass
class RecordTable:
    """Bảng record chuẩn hoá dùng chung cho mọi nguồn (CSV/XLSX/Google Sheets).

    Bất biến (Req 1.10, 1.11):
    - ``columns`` chứa các tên cột DUY NHẤT, theo thứ tự xuất hiện.
    - Mỗi phần tử ``rows`` là dict chỉ chứa khoá thuộc ``columns``.
    """

    columns: List[str] = field(default_factory=list)
    rows: List[Dict[str, str]] = field(default_factory=list)


class DataSourceError(Exception):
    """Lỗi khi đọc/chuẩn hoá nguồn dữ liệu.

    Attributes:
        code: Mã lỗi máy đọc được, ví dụ ``'EMPTY'``, ``'NO_HEADER'``,
            ``'DETECT_FAILED'``, ``'GSHEET_FORBIDDEN'``.
        message: Thông báo tiếng Việt mô tả nguyên nhân để hiển thị cho người dùng.
    """

    def __init__(self, code: str, message: str) -> None:
        self.code = code
        self.message = message
        super().__init__(f"[{code}] {message}")


# ─── Nhận diện encoding ──────────────────────────────────────────────────────


def detect_encoding(raw: bytes) -> str:
    """Nhận diện encoding của dữ liệu thô trong {utf-8, utf-8-sig, windows-1258}.

    Thử lần lượt: UTF-8 BOM (nếu có BOM) → UTF-8 strict → Windows-1258, và trả về
    tên encoding ĐẦU TIÊN giải mã được toàn bộ nội dung không lỗi (Req 1.6). Ưu
    tiên này bảo đảm tiếng Việt có dấu được giữ nguyên Unicode (Req 1.7).

    Args:
        raw: Nội dung file ở dạng bytes.

    Returns:
        Một trong ``{'utf-8', 'utf-8-sig', 'windows-1258'}`` dùng để giải mã.

    Raises:
        DataSourceError: mã ``DETECT_FAILED`` nếu không encoding nào giải mã sạch.
    """
    # BOM UTF-8 → chắc chắn là utf-8-sig (Req 1.6).
    if raw.startswith(codecs.BOM_UTF8):
        return "utf-8-sig"

    for enc in SUPPORTED_ENCODINGS:
        # Khi không có BOM, utf-8-sig tương đương utf-8; bỏ qua để trả 'utf-8'.
        if enc == "utf-8-sig":
            continue
        try:
            raw.decode(enc)
            return enc
        except (UnicodeDecodeError, LookupError):
            continue

    raise DataSourceError(
        "DETECT_FAILED",
        "Không nhận diện được bảng mã (encoding) của file. "
        "Hỗ trợ UTF-8, UTF-8 BOM và Windows-1258.",
    )


# ─── Nhận diện delimiter ─────────────────────────────────────────────────────


def detect_delimiter(header_line: str) -> str:
    """Nhận diện delimiter của dòng tiêu đề trong {',', ';', '\\t'} (Req 1.5).

    Chọn delimiter cho số cột LỚN NHẤT trên dòng tiêu đề. Nếu nhiều delimiter
    cùng cho số cột lớn nhất (> 1) → hoà → không xác định được. Nếu mọi ứng viên
    đều chỉ cho 1 cột → không có delimiter. Cả hai trường hợp ném
    ``DETECT_FAILED`` (Req 1.8).

    Args:
        header_line: Dòng tiêu đề (chuỗi văn bản đã giải mã).

    Returns:
        Một trong ``{',', ';', '\\t'}``.

    Raises:
        DataSourceError: mã ``DETECT_FAILED`` nếu hoà hoặc không tách được cột.
    """
    counts = {d: header_line.count(d) + 1 for d in SUPPORTED_DELIMITERS}
    max_cols = max(counts.values())

    if max_cols <= 1:
        raise DataSourceError(
            "DETECT_FAILED",
            "Không nhận diện được ký tự phân tách cột (dấu phẩy, chấm phẩy hoặc tab) "
            "trên dòng tiêu đề.",
        )

    winners = [d for d, n in counts.items() if n == max_cols]
    if len(winners) != 1:
        raise DataSourceError(
            "DETECT_FAILED",
            "Có nhiều ký tự phân tách khả dĩ cho cùng số cột; không xác định được "
            "ký tự phân tách trên dòng tiêu đề.",
        )

    return winners[0]


# ─── Khử trùng tên cột ───────────────────────────────────────────────────────


def normalize_columns(header: List[str]) -> List[str]:
    """Chuẩn hoá danh sách tên cột: khử ô rỗng và tên trùng, không mất cột (Req 1.11).

    - Ô tiêu đề rỗng (rỗng hoặc chỉ khoảng trắng) → ``"Cột {i}"`` với ``i`` là chỉ
      số cột bắt đầu từ 1.
    - Tên trùng → thêm hậu tố ``"_2"``, ``"_3"``… theo thứ tự xuất hiện.

    Bảo đảm độ dài kết quả bằng độ dài đầu vào và toàn bộ tên là duy nhất.

    Args:
        header: Danh sách tên cột thô lấy từ dòng tiêu đề.

    Returns:
        Danh sách tên cột cùng độ dài, mọi tên duy nhất, theo thứ tự gốc.
    """
    result: List[str] = []
    used: set[str] = set()

    for idx, raw_name in enumerate(header, start=1):
        name = (raw_name or "").strip()
        if not name:
            name = f"Cột {idx}"

        if name in used:
            suffix = 2
            candidate = f"{name}_{suffix}"
            while candidate in used:
                suffix += 1
                candidate = f"{name}_{suffix}"
            name = candidate

        used.add(name)
        result.append(name)

    return result


# ─── Phân tích văn bản phân tách (CSV / TSV / dấu chấm phẩy) ──────────────────


def _is_blank_line(row: List[str]) -> bool:
    """Trả True khi ``row`` ứng với một dòng vật lý HOÀN TOÀN rỗng (Req 1.10).

    ``csv.reader`` trả về danh sách rỗng ``[]`` cho mỗi dòng trống (không chứa ký
    tự nào). Một dòng chỉ gồm các ô rỗng do dấu phân tách (ví dụ ``",,"``) KHÔNG
    bị coi là rỗng — đó là một record hợp lệ với các giá trị rỗng.
    """
    return len(row) == 0


def parse_delimited(
    text: str, delimiter: str, has_header: bool = True
) -> RecordTable:
    """Phân tích văn bản phân tách thành ``RecordTable`` (Req 1.1, 1.10).

    Quy tắc:
    - Bỏ qua mọi dòng hoàn toàn rỗng (Req 1.10).
    - Khi ``has_header`` True: dòng KHÔNG rỗng ĐẦU TIÊN là dòng tiêu đề; tên cột
      được khử trùng/ô rỗng qua :func:`normalize_columns` (Req 1.11).
    - Khi ``has_header`` False: sinh tên cột ``"Cột {i}"`` theo số cột lớn nhất
      của các dòng dữ liệu; mọi dòng không rỗng đều là dữ liệu.
    - Mỗi dòng dữ liệu được ánh xạ thành ``dict`` theo thứ tự cột; ô thiếu → chuỗi
      rỗng, ô dư (vượt số cột) bị bỏ.

    Args:
        text: Nội dung văn bản đã giải mã.
        delimiter: Ký tự phân tách cột (một trong :data:`SUPPORTED_DELIMITERS`).
        has_header: Có lấy dòng đầu làm tiêu đề hay không.

    Returns:
        ``RecordTable`` với cột duy nhất và các dòng dạng ``dict``.

    Raises:
        DataSourceError: mã ``NO_HEADER`` nếu không có dòng KHÔNG rỗng nào.
    """
    reader = csv.reader(io.StringIO(text), delimiter=delimiter)
    # Bỏ các dòng hoàn toàn rỗng, giữ thứ tự xuất hiện (Req 1.10).
    rows: List[List[str]] = [row for row in reader if not _is_blank_line(row)]

    return _rows_to_table(rows, has_header)


def _rows_to_table(rows: List[List[str]], has_header: bool) -> RecordTable:
    """Dựng ``RecordTable`` từ danh sách dòng ĐÃ lọc dòng rỗng (Req 1.10, 1.11).

    Đây là lõi cấu trúc dùng CHUNG cho mọi định dạng (CSV/XLSX): chọn dòng tiêu
    đề, khử trùng tên cột và ánh xạ mỗi dòng dữ liệu thành ``dict`` đồng nhất.
    Việc dùng chung hàm này bảo đảm cùng một bảng nguồn biểu diễn dưới CSV hay
    XLSX cho ra ``RecordTable`` bằng nhau (Req 1.10 / Property 8).

    Args:
        rows: Danh sách dòng (mỗi dòng là list ô chuỗi), ĐÃ loại các dòng hoàn
            toàn rỗng theo quy ước riêng của từng định dạng.
        has_header: Có lấy dòng KHÔNG rỗng đầu tiên làm tiêu đề hay không.

    Returns:
        ``RecordTable`` với cột duy nhất và các dòng dạng ``dict``.

    Raises:
        DataSourceError: mã ``NO_HEADER`` nếu không còn dòng nào sau khi lọc.
    """
    if not rows:
        raise DataSourceError(
            "NO_HEADER",
            "Nguồn dữ liệu không có dòng tiêu đề (chỉ chứa các dòng trống).",
        )

    if has_header:
        header = rows[0]
        data_rows = rows[1:]
    else:
        width = max(len(r) for r in rows)
        header = [""] * width
        data_rows = rows

    columns = normalize_columns(header)

    table = RecordTable(columns=columns, rows=[])
    for raw_row in data_rows:
        record: Dict[str, str] = {}
        for idx, col in enumerate(columns):
            record[col] = raw_row[idx] if idx < len(raw_row) else ""
        table.rows.append(record)

    return table


# ─── Đọc Excel (.xlsx) qua openpyxl ──────────────────────────────────────────


def _cell_to_str(value: object) -> str:
    """Chuẩn hoá giá trị ô Excel về chuỗi (Req 1.2, 1.7).

    - ``None`` (ô trống / ô bị che trong vùng gộp) → chuỗi rỗng.
    - Chuỗi → giữ nguyên (bảo toàn ký tự tiếng Việt Unicode).
    - Số/giá trị khác → ``str(value)``; số nguyên dạng float (vd ``5.0``) được
      rút gọn về ``"5"`` để khớp trực giác người dùng.
    """
    if value is None:
        return ""
    if isinstance(value, str):
        return value
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return str(value)


def _is_blank_xlsx_row(cells: List[str]) -> bool:
    """True khi mọi ô trên dòng đều rỗng sau khi chuẩn hoá (Req 1.10).

    Khác với CSV (dòng vật lý rỗng), một dòng Excel được coi là rỗng khi KHÔNG ô
    nào chứa nội dung — phù hợp ngữ nghĩa "dòng hoàn toàn rỗng" của Req 1.10.
    """
    return all(c == "" for c in cells)


def list_xlsx_sheets(data: bytes) -> List[str]:
    """Trả về danh sách tên sheet của workbook ``.xlsx`` theo thứ tự (Req 1.3).

    Args:
        data: Nội dung file ``.xlsx`` ở dạng bytes.

    Returns:
        Danh sách tên sheet, giữ đúng thứ tự trong workbook.

    Raises:
        DataSourceError: mã ``EMPTY`` nếu ``data`` rỗng; mã ``XLSX_INVALID`` nếu
            không mở được như một workbook hợp lệ.
    """
    if not data:
        raise DataSourceError("EMPTY", "File rỗng, không có dữ liệu để đọc.")

    # Import cục bộ để giữ module thuần (không buộc openpyxl khi chỉ dùng CSV).
    import openpyxl  # noqa: PLC0415

    try:
        wb = openpyxl.load_workbook(
            io.BytesIO(data), read_only=True, data_only=True
        )
    except Exception as exc:  # openpyxl ném nhiều loại lỗi khác nhau
        raise DataSourceError(
            "XLSX_INVALID",
            "Không đọc được file Excel (.xlsx); file có thể hỏng hoặc sai định dạng.",
        ) from exc

    try:
        return list(wb.sheetnames)
    finally:
        wb.close()


def read_xlsx(data: bytes, sheet: Optional[str] = None) -> RecordTable:
    """Đọc file Excel ``.xlsx`` → ``RecordTable`` (Req 1.2, 1.3, 1.12).

    Dùng ``openpyxl`` ở chế độ ``read_only`` cho hiệu năng/bộ nhớ. Quy tắc khớp
    với đường CSV (dùng chung :func:`_rows_to_table`): dòng tiêu đề là dòng KHÔNG
    rỗng đầu tiên, các dòng hoàn toàn rỗng bị bỏ qua, tên cột được khử trùng/ô
    rỗng (Req 1.10, 1.11). Với ô gộp (merged cell), file ``.xlsx`` chỉ lưu giá
    trị tại ô trên-trái còn các ô bị che là rỗng, nên ở chế độ read_only ta đọc
    đúng hành vi "gán về ô trên-trái, các ô khác để rỗng" (Req 1.12).

    Args:
        data: Nội dung file ``.xlsx`` ở dạng bytes.
        sheet: Tên sheet cần đọc; ``None`` → dùng sheet đang active của workbook.

    Returns:
        ``RecordTable`` chuẩn hoá, đồng nhất với đường CSV (Req 1.10).

    Raises:
        DataSourceError: ``EMPTY`` (file rỗng), ``XLSX_INVALID`` (không mở được),
            ``SHEET_NOT_FOUND`` (sheet không tồn tại), ``NO_HEADER`` (chỉ có dòng
            trống / không có dữ liệu).
    """
    if not data:
        raise DataSourceError("EMPTY", "File rỗng, không có dữ liệu để đọc.")

    import openpyxl  # noqa: PLC0415

    try:
        wb = openpyxl.load_workbook(
            io.BytesIO(data), read_only=True, data_only=True
        )
    except Exception as exc:
        raise DataSourceError(
            "XLSX_INVALID",
            "Không đọc được file Excel (.xlsx); file có thể hỏng hoặc sai định dạng.",
        ) from exc

    try:
        if sheet is not None:
            if sheet not in wb.sheetnames:
                raise DataSourceError(
                    "SHEET_NOT_FOUND",
                    f"Không tìm thấy sheet '{sheet}' trong file Excel.",
                )
            ws = wb[sheet]
        else:
            ws = wb.active

        rows: List[List[str]] = []
        for raw_row in ws.iter_rows(values_only=True):
            cells = [_cell_to_str(v) for v in raw_row]
            if _is_blank_xlsx_row(cells):
                continue
            rows.append(cells)
    finally:
        wb.close()

    return _rows_to_table(rows, has_header=True)


# ─── Tải Google Sheets qua URL export CSV (httpx) ────────────────────────────

# Regex bắt spreadsheet id trong các dạng link Google Sheets phổ biến:
#   https://docs.google.com/spreadsheets/d/<ID>/edit#gid=0
#   https://docs.google.com/spreadsheets/d/<ID>/edit?usp=sharing
#   https://docs.google.com/spreadsheets/d/<ID>/export?format=csv&gid=123
#   https://docs.google.com/spreadsheets/d/<ID>
# ID gồm chữ-số, dấu gạch dưới và gạch ngang.
_GSHEET_ID_RE = re.compile(r"/spreadsheets/d/([a-zA-Z0-9_-]+)")

# Thời gian chờ mạng tối đa (giây) cho mỗi yêu cầu tải.
_GSHEET_TIMEOUT_SECONDS: float = 30.0

# Host của trang đăng nhập Google; bị redirect tới đây nghĩa là sheet không công khai.
_GOOGLE_LOGIN_HOSTS: tuple[str, ...] = ("accounts.google.com",)


def _parse_gsheet_url(url: str) -> tuple[str, Optional[str]]:
    """Tách spreadsheet id và ``gid`` (nếu có) từ một link Google Sheets.

    Hỗ trợ ``gid`` nằm ở query (``?gid=123``) hoặc fragment (``#gid=123``).

    Args:
        url: Link Google Sheets do người dùng cung cấp.

    Returns:
        Cặp ``(spreadsheet_id, gid)``; ``gid`` là ``None`` khi link không nêu.

    Raises:
        DataSourceError: mã ``GSHEET_FORBIDDEN`` nếu không tách được spreadsheet id
            (link không phải Google Sheets hợp lệ).
    """
    match = _GSHEET_ID_RE.search(url or "")
    if not match:
        raise DataSourceError(
            "GSHEET_FORBIDDEN",
            "Link Google Sheets không hợp lệ; không tìm thấy mã bảng tính trong đường dẫn.",
        )
    spreadsheet_id = match.group(1)

    parsed = urlparse(url)
    gid: Optional[str] = None
    # gid có thể ở query string...
    query_gid = parse_qs(parsed.query).get("gid")
    if query_gid:
        gid = query_gid[0]
    # ...hoặc ở fragment (#gid=123 hoặc #gid=123&range=...).
    if gid is None and parsed.fragment:
        frag_gid = parse_qs(parsed.fragment).get("gid")
        if frag_gid:
            gid = frag_gid[0]

    return spreadsheet_id, gid


def _build_gsheet_export_url(spreadsheet_id: str, gid: Optional[str]) -> str:
    """Dựng URL export CSV công khai của Google Sheets (Req 1.4)."""
    url = (
        f"https://docs.google.com/spreadsheets/d/{spreadsheet_id}/export?format=csv"
    )
    if gid is not None:
        url += f"&gid={gid}"
    return url


def fetch_gsheet_csv(url: str) -> bytes:
    """Tải dữ liệu một Google Sheets công khai dưới dạng CSV bytes (Req 1.4, 1.13).

    Chuyển link Google Sheets thành URL export CSV rồi tải bằng ``httpx``. Khi
    sheet không được chia sẻ công khai, Google trả HTTP 4xx hoặc redirect tới
    trang đăng nhập (``accounts.google.com``); cả hai trường hợp được quy về lỗi
    ``GSHEET_FORBIDDEN`` thay vì trả về trang HTML đăng nhập (Req 1.13).

    Args:
        url: Link Google Sheets (dạng ``/edit``, có/không ``gid``, hoặc link export).

    Returns:
        Nội dung CSV ở dạng bytes, sẵn sàng đưa vào :func:`_read_csv`.

    Raises:
        DataSourceError: mã ``GSHEET_FORBIDDEN`` khi thiếu quyền truy cập, redirect
            đăng nhập, link sai, hoặc không kết nối được tới Google Sheets.
    """
    import httpx  # noqa: PLC0415 - import cục bộ để giữ module nhẹ khi chỉ dùng CSV/XLSX

    spreadsheet_id, gid = _parse_gsheet_url(url)
    export_url = _build_gsheet_export_url(spreadsheet_id, gid)

    try:
        # follow_redirects=True để theo redirect tải file hợp lệ của Google;
        # ta kiểm host đích để phát hiện redirect tới trang đăng nhập.
        response = httpx.get(
            export_url,
            follow_redirects=True,
            timeout=_GSHEET_TIMEOUT_SECONDS,
        )
    except httpx.HTTPError as exc:
        raise DataSourceError(
            "GSHEET_FORBIDDEN",
            "Không kết nối được tới Google Sheets; kiểm tra lại đường dẫn và kết nối mạng.",
        ) from exc

    # Bị đưa về trang đăng nhập → sheet không công khai (Req 1.13).
    final_host = (response.url.host or "").lower()
    if any(login_host in final_host for login_host in _GOOGLE_LOGIN_HOSTS):
        raise DataSourceError(
            "GSHEET_FORBIDDEN",
            "Google Sheets này không được chia sẻ công khai; vui lòng bật quyền "
            "xem cho 'bất kỳ ai có đường liên kết' rồi thử lại.",
        )

    # HTTP 4xx (vd 401/403/404) → thiếu quyền hoặc không tồn tại (Req 1.13).
    if 400 <= response.status_code < 500:
        raise DataSourceError(
            "GSHEET_FORBIDDEN",
            "Không truy cập được Google Sheets (thiếu quyền hoặc không tồn tại); "
            "vui lòng bật chia sẻ công khai rồi thử lại.",
        )

    response.raise_for_status()
    return response.content


# ─── Điều phối đọc nguồn theo định dạng ──────────────────────────────────────


def _first_non_empty_line(text: str) -> Optional[str]:
    """Trả về dòng vật lý KHÔNG rỗng đầu tiên (raw, không strip) hoặc None.

    Dùng để nhận diện delimiter trên dòng tiêu đề. Định nghĩa "rỗng" khớp với
    :func:`_is_blank_line` (dòng có chuỗi rỗng), nên dòng tiêu đề chọn ở đây
    trùng với dòng tiêu đề mà :func:`parse_delimited` dùng.
    """
    for line in text.splitlines():
        if line != "":
            return line
    return None


def _read_csv(payload: Union[bytes, str], *, has_header: bool = True) -> RecordTable:
    """Đọc nguồn CSV/văn bản phân tách → ``RecordTable`` (Req 1.1, 1.9, 1.10).

    - ``bytes`` → tự nhận diện encoding rồi giải mã (Req 1.6, 1.7).
    - Tự nhận diện delimiter trên dòng tiêu đề (Req 1.5).
    - Nguồn rỗng → ``EMPTY``; chỉ có dòng trống → ``NO_HEADER`` (Req 1.9), KHÔNG
      tạo bảng record.
    """
    if isinstance(payload, bytes):
        if not payload:
            raise DataSourceError("EMPTY", "File rỗng, không có dữ liệu để đọc.")
        encoding = detect_encoding(payload)
        text = payload.decode(encoding)
    elif isinstance(payload, str):
        text = payload
    else:  # pragma: no cover - bảo vệ kiểu đầu vào
        raise DataSourceError(
            "EMPTY", "Không có dữ liệu nguồn hợp lệ để đọc (kiểu không hỗ trợ)."
        )

    if text == "":
        raise DataSourceError("EMPTY", "File rỗng, không có dữ liệu để đọc.")

    header_line = _first_non_empty_line(text)
    if header_line is None:
        raise DataSourceError(
            "NO_HEADER",
            "Nguồn dữ liệu không có dòng tiêu đề (chỉ chứa các dòng trống).",
        )

    delimiter = detect_delimiter(header_line)
    return parse_delimited(text, delimiter, has_header)


def read_source(
    kind: str,
    payload: Union[bytes, str],
    *,
    sheet: Optional[str] = None,
    has_header: bool = True,
) -> RecordTable:
    """Điều phối đọc nguồn dữ liệu theo định dạng và trả về ``RecordTable``.

    Hiện thực đường CSV (Req 1.1, 1.9, 1.10), Excel ``.xlsx`` (Req 1.2, 1.3,
    1.12) và Google Sheets qua URL export CSV (Req 1.4, 1.13).

    Args:
        kind: Loại nguồn — ``'csv'`` | ``'xlsx'`` | ``'gsheet'``.
        payload: Nội dung nguồn (``bytes`` cho file; ``str`` cho văn bản/URL).
        sheet: Tên sheet (chỉ dùng cho ``xlsx``).
        has_header: Dòng đầu là tiêu đề hay không (mặc định True).

    Returns:
        ``RecordTable`` chuẩn hoá dùng chung mọi định dạng (Req 1.10).

    Raises:
        DataSourceError: với mã tương ứng khi nguồn thiếu dữ liệu/tiêu đề, không
            nhận diện được delimiter/encoding, hoặc định dạng chưa hỗ trợ.
    """
    normalized_kind = (kind or "").strip().lower()

    if normalized_kind == "csv":
        return _read_csv(payload, has_header=has_header)

    if normalized_kind in ("xlsx", "excel"):
        if not isinstance(payload, bytes):
            raise DataSourceError(
                "EMPTY",
                "Nguồn Excel (.xlsx) phải là dữ liệu nhị phân (bytes) của file.",
            )
        return read_xlsx(payload, sheet)

    if normalized_kind in ("gsheet", "gsheets", "google-sheets"):
        if not isinstance(payload, str):
            raise DataSourceError(
                "GSHEET_FORBIDDEN",
                "Nguồn Google Sheets phải là một đường liên kết (URL) dạng chuỗi.",
            )
        csv_bytes = fetch_gsheet_csv(payload)
        return _read_csv(csv_bytes, has_header=has_header)

    raise DataSourceError(
        "UNSUPPORTED",
        f"Định dạng nguồn dữ liệu không được hỗ trợ: '{kind}'.",
    )
