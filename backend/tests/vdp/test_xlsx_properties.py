"""Property-based tests cho đường XLSX của Data_Source_Reader.

Targets: ``app.workers.vdp_datasource`` — ``read_xlsx`` và ``list_xlsx_sheets``.

Mỗi property test:
- xây workbook ``.xlsx`` trong bộ nhớ bằng ``openpyxl`` (có sẵn trong venv backend),
- chạy ≥ 100 ví dụ (``@settings(max_examples=100)``, ``deadline=None`` vì có I/O),
- gắn comment tham chiếu theo định dạng
  ``# Feature: vdp-upgrade, Property {n}: {text}``.
"""
import io
import os
import sys

import openpyxl
from hypothesis import given, settings, strategies as st

sys.path.insert(0, os.path.join(os.path.dirname(__file__), '..', '..'))

from app.workers.vdp_datasource import list_xlsx_sheets, read_xlsx


# ─── Sinh nội dung ô (gồm tiếng Việt có dấu) ─────────────────────────────────

_VIETNAMESE_CHARS = (
    "àáảãạăằắẳẵặâầấẩẫậèéẻẽẹêềếểễệ"
    "ìíỉĩịòóỏõọôồốổỗộơờớởỡợ"
    "ùúủũụưừứửữựỳýỷỹỵđ"
    "ÀÁẢÃẠĂẰẮẲẴẶÂẦẤẨẪẬÈÉẺẼẸÊỀẾỂỄỆ"
    "ÌÍỈĨỊÒÓỎÕỌÔỒỐỔỖỘƠỜỚỞỠỢ"
    "ÙÚỦŨỤƯỪỨỬỮỰỲÝỶỸỴĐ"
)


def _cell_text():
    """Nội dung ô: chữ cái/số/khoảng trắng ASCII + tiếng Việt có dấu.

    Loại bỏ ký tự điều khiển (xuống dòng) để tránh nhập nhằng và để openpyxl
    không tự lọc; nội dung là chuỗi để so sánh trực tiếp với ``_cell_to_str``.
    """
    alphabet = st.one_of(
        st.sampled_from(_VIETNAMESE_CHARS),
        st.characters(min_codepoint=0x20, max_codepoint=0x7E),
    )
    # min_size=1: tránh ô rỗng làm dòng/đầu bảng bị coi là trống → khó so khớp.
    return st.text(alphabet=alphabet, min_size=1, max_size=12)


@st.composite
def _xlsx_dataset(draw):
    """Sinh header (tên cột DUY NHẤT, không rỗng) + các dòng đầy đủ N ô.

    Tên cột được thêm tiền tố ``c{i}_`` để bảo đảm duy nhất, tránh phụ thuộc vào
    hành vi khử trùng của parser khi so khớp trực tiếp.
    """
    n_cols = draw(st.integers(min_value=1, max_value=5))
    base = draw(st.lists(_cell_text(), min_size=n_cols, max_size=n_cols))
    header = [f"c{i}_{base[i].strip() or 'col'}" for i in range(n_cols)]

    n_rows = draw(st.integers(min_value=0, max_value=6))
    rows = draw(
        st.lists(
            st.lists(_cell_text(), min_size=n_cols, max_size=n_cols),
            min_size=n_rows,
            max_size=n_rows,
        )
    )
    return header, rows


def _set_text(ws, row, col, value):
    """Ghi ``value`` vào ô (row, col) dưới dạng VĂN BẢN tường minh.

    openpyxl mặc định coi chuỗi bắt đầu bằng ``=`` là công thức (data_type 'f'),
    khi đọc lại với ``data_only=True`` sẽ trả None vì không có giá trị đã tính.
    Đây là hành vi phía GHI workbook, không liên quan tới đúng/sai của
    ``read_xlsx``. Ép ``data_type='s'`` để mọi nội dung được lưu như chuỗi, giúp
    test tập trung vào hành vi ĐỌC.
    """
    cell = ws.cell(row=row, column=col)
    cell.value = value
    cell.data_type = "s"
    return cell


def _build_xlsx(sheets):
    """Dựng workbook ``.xlsx`` trong bộ nhớ.

    Args:
        sheets: danh sách ``(sheet_name, header, rows)``. Sheet đầu tiên là sheet
            active mặc định.

    Returns:
        bytes nội dung file ``.xlsx``.
    """
    wb = openpyxl.Workbook()
    # Xoá sheet mặc định để kiểm soát hoàn toàn tên & thứ tự.
    default = wb.active
    wb.remove(default)

    for name, header, rows in sheets:
        ws = wb.create_sheet(title=name)
        for c, val in enumerate(header, start=1):
            _set_text(ws, 1, c, val)
        for r, row in enumerate(rows, start=2):
            for c, val in enumerate(row, start=1):
                _set_text(ws, r, c, val)

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


# Feature: vdp-upgrade, Property 2: XLSX round-trip bảo toàn dữ liệu
@settings(max_examples=100, deadline=None)
@given(data=_xlsx_dataset())
def test_xlsx_roundtrip_preserves_data(data):
    """Ghi header + rows ra workbook ``.xlsx`` rồi ``read_xlsx`` lại phải phục hồi
    đúng danh sách cột và giá trị từng ô (kể cả tiếng Việt có dấu).

    Tên cột duy nhất nên ``normalize_columns`` giữ nguyên, cho phép so khớp trực
    tiếp cột và giá trị mà không bị nhiễu bởi khử trùng.

    **Validates: Requirements 1.2**
    """
    header, rows = data
    payload = _build_xlsx([("Sheet1", header, rows)])

    table = read_xlsx(payload)

    # Cùng số cột và đúng tên cột (đã duy nhất → giữ nguyên).
    assert table.columns == header

    # Cùng số dòng dữ liệu.
    assert len(table.rows) == len(rows)

    # Mỗi ô được bảo toàn.
    for original_row, parsed_row in zip(rows, table.rows):
        for col, value in zip(header, original_row):
            assert parsed_row[col] == value


# Feature: vdp-upgrade, Property 3: Liệt kê và chọn sheet trong XLSX
@settings(max_examples=100, deadline=None)
@given(
    sheet_names=st.lists(
        st.text(
            alphabet=st.characters(min_codepoint=0x41, max_codepoint=0x5A),
            min_size=1,
            max_size=6,
        ),
        min_size=2,
        max_size=4,
        unique=True,
    ),
    payloads=st.lists(_xlsx_dataset(), min_size=2, max_size=4),
)
def test_xlsx_list_and_select_sheets(sheet_names, payloads):
    """Với workbook nhiều sheet tên phân biệt: ``list_xlsx_sheets`` trả đúng tập
    tên theo thứ tự, và ``read_xlsx(data, sheet=name)`` đọc đúng dữ liệu của sheet
    được chọn (dữ liệu mỗi sheet khác nhau để xác nhận đúng sheet).

    **Validates: Requirements 1.3**
    """
    # Ghép số lượng sheet = min(len) để mỗi sheet có dataset riêng biệt.
    n = min(len(sheet_names), len(payloads))
    names = sheet_names[:n]

    sheets = []
    for i in range(n):
        header, rows = payloads[i]
        # Làm dữ liệu mỗi sheet khác biệt rõ ràng bằng tiền tố theo tên sheet.
        header = [f"{names[i]}_{h}" for h in header]
        sheets.append((names[i], header, rows))

    payload = _build_xlsx(sheets)

    # list_xlsx_sheets trả đúng tên theo thứ tự khai báo.
    assert list_xlsx_sheets(payload) == names

    # read_xlsx với từng sheet đọc đúng dữ liệu của chính sheet đó.
    for name, header, rows in sheets:
        table = read_xlsx(payload, sheet=name)
        assert table.columns == header
        assert len(table.rows) == len(rows)
        for original_row, parsed_row in zip(rows, table.rows):
            for col, value in zip(header, original_row):
                assert parsed_row[col] == value


# Feature: vdp-upgrade, Property 4: Merged cell gán về ô trên-trái
@settings(max_examples=100, deadline=None)
@given(
    n_cols=st.integers(min_value=2, max_value=5),
    n_data_rows=st.integers(min_value=2, max_value=5),
    value=_cell_text(),
    orientation=st.sampled_from(["horizontal", "vertical"]),
)
def test_xlsx_merged_cell_top_left(n_cols, n_data_rows, value, orientation):
    """Khi tạo vùng merge (ngang hoặc dọc) với một giá trị, ``read_xlsx`` gán giá
    trị về ô trên-trái còn các ô bị che trong vùng để rỗng.

    **Validates: Requirements 1.12**
    """
    wb = openpyxl.Workbook()
    ws = wb.active

    # Hàng 1 là tiêu đề (KHÔNG rỗng) để có header hợp lệ; tên cột duy nhất.
    header = [f"H{i}" for i in range(n_cols)]
    for c, val in enumerate(header, start=1):
        _set_text(ws, 1, c, val)

    # Điền dữ liệu nền duy nhất cho mọi ô để các dòng không bị coi là rỗng.
    # Dùng "x{r}_{c}" để mỗi ô khác nhau và khác value gộp.
    for r in range(n_data_rows):
        for c in range(n_cols):
            _set_text(ws, r + 2, c + 1, f"x{r}_{c}")

    # Chọn vùng merge bắt đầu tại ô dữ liệu trên-trái: hàng excel 2, cột 1.
    # openpyxl dùng chỉ số 1-based; hàng 1 = header → dữ liệu bắt đầu hàng 2.
    start_row = 2
    start_col = 1
    if orientation == "horizontal":
        end_row = start_row
        end_col = start_col + 1  # gộp 2 ô ngang
    else:
        end_row = start_row + 1  # gộp 2 ô dọc
        end_col = start_col

    ws.merge_cells(
        start_row=start_row,
        start_column=start_col,
        end_row=end_row,
        end_column=end_col,
    )
    # Gán giá trị vào ô trên-trái của vùng merge (dạng văn bản tường minh).
    _set_text(ws, start_row, start_col, value)

    buf = io.BytesIO()
    wb.save(buf)
    payload = buf.getvalue()

    table = read_xlsx(payload)

    # Ánh xạ vị trí merge (1-based, gồm header) → chỉ số trong RecordTable.
    # RecordTable: header bị tách ra, nên data row index = excel_row - 2.
    columns = table.columns

    tl_row_idx = start_row - 2
    tl_col_idx = start_col - 1
    # Ô trên-trái giữ đúng giá trị đã gán.
    assert table.rows[tl_row_idx][columns[tl_col_idx]] == value

    # Các ô còn lại trong vùng merge phải rỗng.
    for er in range(start_row, end_row + 1):
        for ec in range(start_col, end_col + 1):
            if er == start_row and ec == start_col:
                continue
            row_idx = er - 2
            col_idx = ec - 1
            assert table.rows[row_idx][columns[col_idx]] == ""
